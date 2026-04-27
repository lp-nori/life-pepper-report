"""
python-pptx ベースのプレゼンテーション生成モジュール
LIFE PEPPER ブランドガイドライン（v4テンプレート定数準拠）

スライド構成:
  1. 表紙
  2. 数字サマリー
  3. 最重要課題①
  4. 最重要課題②
  5. 最重要課題③
  6. 今週の意思決定・アクション
  7. 注視案件（契約リスク）
  8. 先週との変化
"""
from __future__ import annotations
import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree


# ─── ブランド定数（create_template_v4.js 準拠） ──────────────────

C_NAVY   = RGBColor(0x1B, 0x14, 0x64)   # #1B1464
C_ORANGE = RGBColor(0xFF, 0x6B, 0x35)   # #FF6B35
C_BG     = RGBColor(0xF7, 0xF7, 0xFB)   # #F7F7FB
C_CARD   = RGBColor(0xFF, 0xFF, 0xFF)
C_GRAY   = RGBColor(0xE8, 0xE8, 0xF0)   # #E8E8F0
C_DARK   = RGBColor(0x1A, 0x1A, 0x2E)   # #1A1A2E
C_MID    = RGBColor(0x7A, 0x7A, 0x92)   # #7A7A92
C_LIGHT  = RGBColor(0xA0, 0xA0, 0xB4)   # #A0A0B4
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_GREEN  = RGBColor(0x27, 0xAE, 0x60)
C_RED    = RGBColor(0xE7, 0x4C, 0x3C)

FONT_JP = "Noto Sans JP"
FONT_EN = "Noto Sans"

# レイアウト定数（4:3、単位: インチ）
SW = 10.0    # スライド幅
SH = 7.5     # スライド高さ
M  = 0.66    # 左右余白

TAG_Y     = 0.28   # セクションタグ Y座標
TITLE_Y   = 0.77   # スライドタイトル Y座標
CONTENT_Y = 1.65   # コンテンツ開始 Y座標
FY        = 6.95   # フッター Y座標

FOOTER_TEXT = (
    "\u00A9 \u682A\u5F0F\u4F1A\u793E LIFE PEPPER (LIFE PEPPER INC) "
    "All Rights Reserved.  |  CONFIDENTIAL"
)


# ─── 基本描画ヘルパー ─────────────────────────────────────────────

def _i(v: float):  return Inches(v)
def _pt(v: float): return Pt(v)


def _blank_layout(prs: Presentation):
    for layout in prs.slide_layouts:
        if layout.name in ("Blank", "空白"):
            return layout
    return prs.slide_layouts[-1]


def _set_bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, x: float, y: float, w: float, h: float,
          color: RGBColor, border: bool = False):
    sp = slide.shapes.add_shape(1, _i(x), _i(y), _i(w), _i(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if border:
        sp.line.color.rgb = C_GRAY
        sp.line.width = _pt(0.4)
    else:
        sp.line.fill.background()
    return sp


def _tb(slide, text: str, x: float, y: float, w: float, h: float, *,
        font: str = None, size: float = 11, bold: bool = False,
        color: RGBColor = None, align: str = "left"):
    txBox = slide.shapes.add_textbox(_i(x), _i(y), _i(w), _i(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }.get(align, PP_ALIGN.LEFT)
    run = p.add_run()
    run.text = text
    run.font.name = font or FONT_JP
    run.font.size = _pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    return txBox


def _footer(slide):
    _rect(slide, M, FY, SW - M * 2, 0.01, C_GRAY)
    _tb(slide, FOOTER_TEXT, M, FY + 0.06, SW - M * 2, 0.22,
        font=FONT_EN, size=6.5, color=C_LIGHT, align="left")


def _section_tag(slide, text: str):
    """オレンジバー＋セクション名ラベル"""
    _rect(slide, M, TAG_Y, 0.044, 0.33, C_ORANGE)
    _tb(slide, text, M + 0.15, TAG_Y - 0.02, 6.5, 0.38,
        size=10, bold=True, color=C_NAVY)


def _slide_title(slide, text: str):
    _tb(slide, text, M, TITLE_Y, 8.0, 0.77, size=24, bold=True, color=C_NAVY)


def _set_cell_fill(cell, hex_color: str):
    """テーブルセル背景色をXMLで直接設定"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        local = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if local in ("solidFill", "gradFill", "noFill", "pattFill"):
            tcPr.remove(child)
    sf = etree.SubElement(tcPr, qn("a:solidFill"))
    sr = etree.SubElement(sf, qn("a:srgbClr"))
    sr.set("val", hex_color)


def _draw_table(slide, rows: list[list[str]],
                x: float, y: float, w: float, max_h: float = 4.8):
    if not rows:
        return
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows)
    row_h = min(max_h / n_rows, 0.58)
    tbl = slide.shapes.add_table(
        n_rows, n_cols,
        _i(x), _i(y), _i(w), _i(row_h * n_rows)
    ).table
    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            cell = tbl.cell(ri, ci)
            text = row[ci] if ci < len(row) else ""
            cell.text = text
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run() if not p.runs else p.runs[0]
            run.font.name = FONT_JP
            run.font.size = _pt(10 if ri == 0 else 9.5)
            run.font.bold = (ri == 0)
            if ri == 0:
                run.font.color.rgb = C_WHITE
                _set_cell_fill(cell, "1B1464")
            elif ri % 2 == 0:
                run.font.color.rgb = C_DARK
                _set_cell_fill(cell, "F4F4F9")
            else:
                run.font.color.rgb = C_DARK


# ─── データ抽出パーサー ──────────────────────────────────────────

def _extract_section(md: str, heading_pat: str) -> str:
    m = re.search(heading_pat + r".*?\n(.*?)(?=\n## |\Z)", md, re.DOTALL)
    return m.group(1).strip() if m else ""


def _parse_table_rows(block: str) -> list[list[str]]:
    rows = []
    for line in block.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        if re.match(r"^\|[\s\-:|]+\|$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if any(cells):
            rows.append(cells)
    return rows


def _parse_issues(block: str) -> list[tuple[str, str, str, str]]:
    """課題ブロックを最大3件の (title, fact, impact, decision) タプルに分解"""
    parts = re.split(r"\n(?=\*\*[①②③]|\d+\.\s*\*\*)", block)
    parts = [p.strip() for p in parts if p.strip()]
    results = []
    for part in parts[:3]:
        m = re.match(r"\*\*[①②③]?\s*(.+?)\*\*", part)
        title = m.group(1).strip() if m else (
            part.splitlines()[0].lstrip("*①②③1234. ").strip()
        )
        fact = impact = decision = ""
        for line in part.splitlines():
            body = re.sub(r"^[^:：]+[:：]\s*", "", line.strip().lstrip("-・").strip())
            if "事実" in line or "課題の" in line:
                fact = body
            elif "影響" in line or "放置" in line:
                impact = body
            elif "意思決定" in line or "必要な" in line:
                decision = body
        # フォールバック：3行箇条書き
        if not any([fact, impact, decision]):
            bullets = [
                l.strip().lstrip("-・*①②③ ").strip()
                for l in part.splitlines()[1:]
                if l.strip().lstrip("-・*").strip()
            ]
            if len(bullets) >= 1: fact     = bullets[0]
            if len(bullets) >= 2: impact   = bullets[1]
            if len(bullets) >= 3: decision = bullets[2]
        results.append((title, fact, impact, decision))
    return results


def _parse_changes(block: str) -> tuple[str, str]:
    good = bad = ""
    for line in block.splitlines():
        l = line.strip().lstrip("-・●*").strip()
        if not l:
            continue
        body = re.sub(r"^[^:：]+[:：]\s*", "", l)
        if any(w in line for w in ("良くな", "改善", "増加", "達成", "↑")):
            if not good: good = body
        elif any(w in line for w in ("悪くな", "低下", "遅延", "未達", "↓")):
            if not bad: bad = body
    lines = [
        l.strip().lstrip("-・*●").strip()
        for l in block.splitlines()
        if l.strip().lstrip("-・*●").strip()
    ]
    if not good and lines:          good = lines[0]
    if not bad and len(lines) > 1:  bad  = lines[1]
    return good, bad


def _parse_mrr(mrr_summary: str) -> tuple[str, str, str]:
    mrr_val = diff_val = judgment = "未集計"
    m = re.search(r"([\d,]+)万円\s*/\s*目標", mrr_summary)
    if m:
        try:
            v = int(m.group(1).replace(",", ""))
            mrr_val  = f"{v:,}"
            d        = v - 2400
            diff_val = f"{'+'if d>=0 else ''}{d:,}"
            judgment = "✅" if v/2400 >= 0.9 else ("⚠️" if v/2400 >= 0.75 else "🚨")
        except ValueError:
            pass
    return mrr_val, diff_val, judgment


# ─── スライドビルダー ─────────────────────────────────────────────

def _slide1_cover(prs, today_str: str, mrr_val: str,
                  diff_display: str, judgment: str):
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_bg(slide, C_BG)
    _rect(slide, 0, 0, 0.083, SH, C_NAVY)        # 左縦バー
    _rect(slide, 0, 5.5, SW, 2.0, C_NAVY)         # 下部帯

    _tb(slide, "MD週次報告",
        M + 0.12, 1.9, 7.5, 1.3, size=44, bold=True, color=C_NAVY)
    _rect(slide, M + 0.12, 3.38, 1.6, 0.05, C_ORANGE)
    _tb(slide, today_str,
        M + 0.12, 3.55, 7.5, 0.55, size=14, color=C_MID)

    mrr_line = f"MRR  {mrr_val}万円    目標比 {diff_display}万円    {judgment}"
    _tb(slide, mrr_line,
        0.9, 5.65, 8.4, 0.66,
        font=FONT_EN, size=17, bold=True, color=C_WHITE)

    _tb(slide, FOOTER_TEXT, 0.44, SH - 0.33, 9.0, 0.22,
        font=FONT_EN, size=6.5, color=RGBColor(0x66, 0x66, 0xAA))


def _slide2_summary(prs, section1: str, mrr_summary: str):
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_bg(slide, C_BG)
    _section_tag(slide, "数字サマリー")
    _slide_title(slide, "数字サマリー")
    rows = _parse_table_rows(section1)
    if rows:
        _draw_table(slide, rows, M, CONTENT_Y, SW - M * 2, 4.5)
    else:
        _tb(slide, mrr_summary, M, CONTENT_Y, SW - M * 2, 4.5,
            size=11, color=C_DARK)
    _footer(slide)


def _slide_issue(prs, label: str, title: str,
                 fact: str, impact: str, decision: str):
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_bg(slide, C_BG)
    _section_tag(slide, label)
    _slide_title(slide, title or "（データなし）")

    items = [
        ("課題の事実",        fact     or "（データなし）", C_NAVY),
        ("放置した場合の影響",  impact   or "（データなし）", C_ORANGE),
        ("必要な意思決定",     decision or "（データなし）", C_GREEN),
    ]
    y = CONTENT_Y
    card_h = 1.38
    for lbl_text, body, accent in items:
        # カード白背景
        sp = slide.shapes.add_shape(1, _i(M), _i(y), _i(SW - M * 2), _i(card_h))
        sp.fill.solid()
        sp.fill.fore_color.rgb = C_CARD
        sp.line.fill.background()
        # 左アクセントバー
        _rect(slide, M, y, 0.055, card_h, accent)
        # ラベル（小）
        _tb(slide, lbl_text, M + 0.2, y + 0.07, 3.5, 0.3,
            size=8.5, bold=True, color=C_MID)
        # 本文
        _tb(slide, body, M + 0.2, y + 0.35, SW - M * 2 - 0.38, 0.9,
            size=11, color=C_DARK)
        y += card_h + 0.13
    _footer(slide)


def _slide6_actions(prs, section3: str):
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_bg(slide, C_BG)
    _section_tag(slide, "今週の意思決定・アクション")
    _slide_title(slide, "今週の意思決定・アクション")
    rows = _parse_table_rows(section3)
    if rows:
        _draw_table(slide, rows, M, CONTENT_Y, SW - M * 2, 4.8)
    else:
        lines = [
            l.strip().lstrip("-・").strip()
            for l in section3.splitlines()
            if l.strip().lstrip("-・").strip()
        ]
        y = CONTENT_Y
        for line in lines[:5]:
            sp = slide.shapes.add_shape(1, _i(M), _i(y), _i(SW - M * 2), _i(0.66))
            sp.fill.solid(); sp.fill.fore_color.rgb = C_CARD; sp.line.fill.background()
            _rect(slide, M, y, 0.055, 0.66, C_NAVY)
            _tb(slide, line, M + 0.2, y + 0.1, SW - M * 2 - 0.38, 0.5,
                size=10.5, color=C_DARK)
            y += 0.75
    _footer(slide)


def _slide7_watchlist(prs, section4: str):
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_bg(slide, C_BG)
    _section_tag(slide, "注視案件（契約リスク）")
    _slide_title(slide, "注視案件（契約リスク）")
    rows = _parse_table_rows(section4)
    if rows:
        _draw_table(slide, rows, M, CONTENT_Y, SW - M * 2, 4.8)
    else:
        _tb(slide, section4 or "注視案件なし",
            M, CONTENT_Y, SW - M * 2, 4.5, size=10.5, color=C_DARK)
    _footer(slide)


def _slide8_changes(prs, section5: str):
    slide = prs.slides.add_slide(_blank_layout(prs))
    _set_bg(slide, C_BG)
    _section_tag(slide, "先週との変化")
    _slide_title(slide, "先週との変化")

    good, bad = _parse_changes(section5)
    card_w = (SW - M * 2 - 0.3) / 2
    card_h = 4.1

    for cx, (emoji, label, text, accent) in [
        (M,                ("✅", "良くなったこと", good, C_GREEN)),
        (M + card_w + 0.3, ("🚨", "悪くなったこと", bad,  C_RED)),
    ]:
        sp = slide.shapes.add_shape(1, _i(cx), _i(CONTENT_Y), _i(card_w), _i(card_h))
        sp.fill.solid(); sp.fill.fore_color.rgb = C_CARD; sp.line.fill.background()
        _rect(slide, cx, CONTENT_Y, card_w, 0.055, accent)
        _tb(slide, f"{emoji}  {label}",
            cx + 0.2, CONTENT_Y + 0.12, card_w - 0.4, 0.44,
            size=12, bold=True, color=accent)
        _tb(slide, text or "（データなし）",
            cx + 0.2, CONTENT_Y + 0.7, card_w - 0.4, card_h - 0.85,
            size=13, color=C_DARK)
    _footer(slide)


# ─── メインエントリポイント ───────────────────────────────────────

def build_report_pptx(report_md: str, mrr_summary: str,
                      today_str: str, output_path: str) -> str:
    """レポートMarkdownから8スライドPPTXを生成して保存する。"""
    prs = Presentation()
    prs.slide_width  = _i(SW)
    prs.slide_height = _i(SH)

    mrr_val, diff_val, judgment = _parse_mrr(mrr_summary)
    try:
        d = int(mrr_val.replace(",", "")) - 2400
        diff_display = f"{'+'if d>=0 else '▲'}{abs(d):,}"
    except ValueError:
        diff_display = diff_val

    s1 = _extract_section(report_md, r"## 1\.")   # 数字サマリー
    s3 = _extract_section(report_md, r"## 3\.")   # 最重要課題（全件）
    s4 = _extract_section(report_md, r"## 4\.")   # 意思決定・アクション
    s5 = _extract_section(report_md, r"## 5\.")   # 注視案件
    s7 = _extract_section(report_md, r"## 7\.")   # 先週との変化

    issues = _parse_issues(s3)

    _slide1_cover(prs, today_str, mrr_val, diff_display, judgment)
    _slide2_summary(prs, s1, mrr_summary)
    for i, lbl in enumerate(["最重要課題①", "最重要課題②", "最重要課題③"]):
        t, f, imp, d = issues[i] if i < len(issues) else ("データなし", "", "", "")
        _slide_issue(prs, lbl, t, f, imp, d)
    _slide6_actions(prs, s4)
    _slide7_watchlist(prs, s5)
    _slide8_changes(prs, s7)

    prs.save(output_path)
    return output_path
